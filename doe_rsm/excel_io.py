"""Excelのプロジェクトファイル：設定の読み込み、実験計画（実験指示書）の書き出し、Yの読み込み、解析結果の書き出し。

計算はすべてPython側で行い、Excelには値だけを書く（セル操作で数式が壊れる心配をなくすため）。
黄色・青字のセルが入力欄。
"""
from datetime import datetime
from typing import Dict, List, Optional

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__
from .design import FACE_CCD, Design
from .factors import Factor, Response
from .model import term_names
from .project import (DESIGN_LABELS, GOAL_LABELS, RANDOMIZATION_LABELS, Project, RSMResult, ScreeningResult)

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FONT = Font(color="1F4E79")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
BOLD = Font(bold=True)
OK_FONT = Font(color="1E7B34")
NG_FONT = Font(color="C00000", bold=True)
WARN_FONT = Font(color="9C5700")

S_PROJECT, S_FACTORS, S_DESIGN, S_HISTORY = "Project", "Factors", "Design", "History"
RESULT_SHEETS = ("Model", "ANOVA", "Gates", "Canonical", "Optimization", "OperatingWindow", "Residuals",
                 "Effects", "SteepestPath")

PROJECT_ROWS = [
    ("テーマ", "title", "実験の目的が分かる名前"),
    ("目的変数", "response_name", "例：剥離強度"),
    ("単位", "response_unit", ""),
    ("最適化の方向", "goal", "最大化→LSLのみ使用 / 最小化→USLのみ / 目標値→LSL・USL・T を使用"),
    ("下限規格 LSL", "lsl", ""),
    ("上限規格 USL", "usl", ""),
    ("目標値 T", "target", ""),
    ("予測区間の信頼水準", "confidence", "通常0.95。最大化・最小化は片側、目標値は両側で判定"),
    ("計画の種類", "design_type", "スクリーニング：Plackett-Burman / 2水準要因　応答曲面：面心CCD / 回転可能CCD / Box-Behnken"),
    ("中心点の数", "n_center", "純誤差・曲率の評価に使うため3以上。RSMでは5〜6を推奨"),
    ("PB実験数", "pb_runs", "Plackett-Burmanのとき。空欄なら因子数から自動（8/12/16/20/24）"),
    ("ランダマイズ方式", "randomization", "グループ化＝変更困難因子の水準ごとにまとめる順序。解析は完全ランダムとみなした近似"),
    ("乱数シード", "seed", "同じシードなら同じ実験順序を再現できる"),
]
FACTOR_HEADERS = ["No", "因子名", "単位", "下限(-1)", "上限(+1)", "安全下限", "安全上限", "変更困難(○)"]


def _inv(d: Dict[str, str]) -> Dict[str, str]:
    return {v: k for k, v in d.items()}


def _num(v) -> Optional[float]:
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    return float(v)


def _style_header(ws, row: int, ncol: int, col0: int = 1):
    for c in range(col0, col0 + ncol):
        cell = ws.cell(row, c)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _mark_input(cell):
    cell.fill = INPUT_FILL
    cell.font = INPUT_FONT


def _widths(ws, widths: List[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fresh_sheet(wb, name: str):
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        del wb[name]
        return wb.create_sheet(name, idx)
    return wb.create_sheet(name)


def _status_font(text: str):
    if text.startswith("×"):
        return NG_FONT
    if text.startswith("△"):
        return WARN_FONT
    if text.startswith("○"):
        return OK_FONT
    return None


# ---------------------------------------------------------------- テンプレート・設定

def write_template(path: str, n_factors: int = 3, title: str = "", project: Optional[Project] = None) -> None:
    """設定を入力するための新しいプロジェクトファイルを作る。project を渡すとその値で埋める。"""
    wb = Workbook()
    ws = wb.active
    ws.title = S_PROJECT
    ws["A1"] = "プロジェクト設定（黄色セルに入力）"
    ws["A1"].font = Font(bold=True, size=13)
    p = project
    values = {} if p is None else {
        "title": p.title, "response_name": p.response.name, "response_unit": p.response.unit,
        "goal": GOAL_LABELS[p.response.goal], "lsl": p.response.lsl, "usl": p.response.usl,
        "target": p.response.target, "confidence": p.response.confidence,
        "design_type": DESIGN_LABELS[p.design_type], "n_center": p.n_center, "pb_runs": p.pb_runs,
        "randomization": RANDOMIZATION_LABELS[p.randomization], "seed": p.seed,
    }
    defaults = {"title": title, "goal": "最大化", "confidence": 0.95, "design_type": DESIGN_LABELS[FACE_CCD],
                "n_center": 6, "randomization": RANDOMIZATION_LABELS["complete"], "seed": 1}
    for i, (label, key, help_) in enumerate(PROJECT_ROWS, start=3):
        ws.cell(i, 1, label).font = BOLD
        c = ws.cell(i, 2, values.get(key, defaults.get(key)))
        _mark_input(c)
        ws.cell(i, 3, help_)
    rows = {key: i for i, (_, key, _) in enumerate(PROJECT_ROWS, start=3)}
    for key, options in (("goal", GOAL_LABELS.values()), ("design_type", DESIGN_LABELS.values()),
                         ("randomization", RANDOMIZATION_LABELS.values())):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(ws.cell(rows[key], 2))
    _widths(ws, [20, 22, 90])

    wf = wb.create_sheet(S_FACTORS)
    wf["A1"] = "因子（範囲は±1＝要因点の実値。安全限界は空欄なら制限なし）"
    wf["A1"].font = Font(bold=True, size=13)
    for c, h in enumerate(FACTOR_HEADERS, start=1):
        wf.cell(3, c, h)
    _style_header(wf, 3, len(FACTOR_HEADERS))
    factors = p.factors if p is not None else [None] * n_factors
    for i, f in enumerate(factors):
        r = 4 + i
        wf.cell(r, 1, i + 1)
        vals = ([f.name, f.unit, f.low, f.high, f.safety_low, f.safety_high, "○" if f.hard_to_change else None]
                if f is not None else [f"因子{i + 1}", "", None, None, None, None, None])
        for j, v in enumerate(vals, start=2):
            _mark_input(wf.cell(r, j, v))
    _widths(wf, [5, 16, 10, 11, 11, 11, 11, 12])

    wh = wb.create_sheet(S_HISTORY)
    for c, h in enumerate(["日時", "操作", "内容", "バージョン"], start=1):
        wh.cell(1, c, h)
    _style_header(wh, 1, 4)
    _widths(wh, [20, 14, 80, 12])
    _append_history(wb, "作成", "プロジェクトファイルを作成")
    wb.save(path)


def _append_history(wb, action: str, detail: str):
    wh = wb[S_HISTORY] if S_HISTORY in wb.sheetnames else wb.create_sheet(S_HISTORY)
    wh.append([datetime.now().strftime("%Y-%m-%d %H:%M"), action, detail, __version__])


def read_project(path: str) -> Project:
    """Project・Factors（あれば Design の Y と Model の使用項も）を読む。"""
    wb = load_workbook(path, data_only=True)
    ws = wb[S_PROJECT]
    raw = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        for lab, key, _ in PROJECT_ROWS:
            if label == lab:
                raw[key] = ws.cell(r, 2).value
    goal = _inv(GOAL_LABELS).get(raw.get("goal"))
    if goal is None:
        raise ValueError(f"最適化の方向が不正: {raw.get('goal')}")
    design_type = _inv(DESIGN_LABELS).get(raw.get("design_type"))
    if design_type is None:
        raise ValueError(f"計画の種類が不正: {raw.get('design_type')}")
    response = Response(raw.get("response_name") or "Y", goal, _num(raw.get("lsl")), _num(raw.get("usl")),
                        _num(raw.get("target")), _num(raw.get("confidence")) or 0.95,
                        raw.get("response_unit") or "")

    wf = wb[S_FACTORS]
    factors = []
    for r in range(4, wf.max_row + 1):
        name = wf.cell(r, 2).value
        if name in (None, ""):
            continue
        low, high = _num(wf.cell(r, 4).value), _num(wf.cell(r, 5).value)
        if low is None or high is None or high <= low:
            raise ValueError(f"因子「{name}」の下限・上限を確認（下限 < 上限）")
        factors.append(Factor(str(name), low, high, wf.cell(r, 3).value or "", _num(wf.cell(r, 6).value),
                              _num(wf.cell(r, 7).value), wf.cell(r, 8).value == "○"))
    if len(factors) < 2:
        raise ValueError("因子は2つ以上必要")

    pb = _num(raw.get("pb_runs"))
    project = Project(raw.get("title") or "", factors, response, design_type, int(_num(raw.get("n_center")) or 0),
                      int(_num(raw.get("seed")) or 1),
                      _inv(RANDOMIZATION_LABELS).get(raw.get("randomization"), "complete"),
                      int(pb) if pb else None)
    if S_DESIGN in wb.sheetnames:
        _read_design(wb[S_DESIGN], project)
    if "Model" in wb.sheetnames and not project.is_screening:
        project.used_terms = _read_used_terms(wb["Model"], len(factors))
    return project


def _read_design(ws, project: Project):
    """Design シートから計画と Y を復元する。計画は設定から作り直し、シートのコード値と一致するか確かめる。"""
    d = project.make_design()
    k = len(project.factors)
    header = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    c_std = header.index("標準順") + 1
    c_order = header.index("実験順序") + 1
    order = np.zeros(d.n_runs, int)
    c_x = header.index("X1(コード)") + 1
    c_y = next(i for i, h in enumerate(header, start=1) if str(h).startswith("Y:"))
    y = np.full(d.n_runs, np.nan)
    seen = 0
    for r in range(4, ws.max_row + 1):
        std = ws.cell(r, c_std).value
        if std is None:
            continue
        i = int(std) - 1
        x = [float(ws.cell(r, c_x + j).value) for j in range(k)]
        if not np.allclose(x, d.points[i]):
            raise ValueError(f"Design シートの標準順{std}のコード値が設定と合わない（因子・計画の種類を変えた？）")
        order[i] = int(ws.cell(r, c_order).value)
        v = _num(ws.cell(r, c_y).value)
        if v is not None:
            y[i] = v
        seen += 1
    if seen != d.n_runs:
        raise ValueError(f"Design シートの行数（{seen}）が計画（{d.n_runs}）と合わない")
    d.run_order = order
    project.y = y


def _read_used_terms(ws, k: int) -> Optional[List[int]]:
    names = term_names(k)
    used = {}
    for r in range(4, ws.max_row + 1):
        name, flag = ws.cell(r, 1).value, ws.cell(r, 2).value
        if name in names and flag in (0, 1):
            used[name] = int(flag)
    if len(used) != len(names):
        return None
    return [used[n] for n in names]


# ---------------------------------------------------------------- 実験計画

def write_design(path: str, project: Project, force: bool = False) -> Design:
    """設定から計画を作り、実験順に並べた Design シート（実験指示書・Y入力欄）を書く。

    既に Y が入っている Design シートは、force=True でない限り上書きしない。
    """
    wb = load_workbook(path)
    if S_DESIGN in wb.sheetnames and not force:
        old = wb[S_DESIGN]
        header = [old.cell(3, c).value for c in range(1, old.max_column + 1)]
        yc = next((i for i, h in enumerate(header, start=1) if str(h).startswith("Y:")), None)
        if yc and any(old.cell(r, yc).value not in (None, "") for r in range(4, old.max_row + 1)):
            raise ValueError("Design シートに Y が入力済み。作り直すなら force=True（--force）")
    d = project.make_design()
    for msg in d.notes:
        _append_history(wb, "注意", msg)
    _write_design_sheet(wb, project,
                        f"実験指示書（{DESIGN_LABELS[project.design_type]}・{d.n_runs}回・乱数シード {project.seed}）：この順に実験し、Yを入力")
    _append_history(wb, "計画", f"{DESIGN_LABELS[project.design_type]} {d.n_runs}回、中心点{project.n_center}、"
                                 f"{RANDOMIZATION_LABELS[project.randomization]}、シード{project.seed}")
    wb.save(path)
    return d


def _write_design_sheet(wb, project: Project, title: str):
    d = project.design
    ws = _fresh_sheet(wb, S_DESIGN)
    k = d.k
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("○ 全実験点が安全限界内" if d.is_safe else
                "× 安全限界を超える実験点あり → この計画は実行不可。範囲を縮小するか計画の種類を変える")
    ws["A2"].font = _status_font(ws["A2"].value)
    r = project.response
    headers = (["実験順序", "標準順", "区分"] + [f"X{i + 1}(コード)" for i in range(k)] + _real_headers(project) +
               [f"Y:{r.name}[{r.unit}]" if r.unit else f"Y:{r.name}", "メモ（異常・気付き）"])
    for c, h in enumerate(headers, start=1):
        ws.cell(3, c, h)
    _style_header(ws, 3, len(headers))
    real = d.real_points()
    y = project.y if project.y is not None else np.full(d.n_runs, np.nan)
    for row_i, std_i in enumerate(np.argsort(d.run_order, kind="stable"), start=4):
        vals = ([int(d.run_order[std_i]), int(std_i + 1), d.kinds[std_i]] + [float(v) for v in d.points[std_i]] +
                [float(v) for v in real[std_i]] + [None if np.isnan(y[std_i]) else float(y[std_i])])
        for c, v in enumerate(vals, start=1):
            ws.cell(row_i, c, v)
        _mark_input(ws.cell(row_i, len(headers) - 1))
        _mark_input(ws.cell(row_i, len(headers)))
    _widths(ws, [9, 7, 8] + [10] * k + [14] * k + [16, 30])
    ws.freeze_panes = "A4"
    if wb.sheetnames.index(S_DESIGN) != 2:
        wb.move_sheet(S_DESIGN, offset=2 - wb.sheetnames.index(S_DESIGN))


# ---------------------------------------------------------------- 解析結果

def _table(ws, top: int, left: int, headers: List[str], rows: List[list], fmt: Optional[Dict[int, str]] = None):
    for c, h in enumerate(headers):
        ws.cell(top, left + c, h)
    _style_header(ws, top, len(headers), left)
    for i, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            if isinstance(v, (np.floating, np.integer)):
                v = v.item()
            if isinstance(v, float) and not np.isfinite(v):
                v = None
            cell = ws.cell(top + i, left + c, v)
            if fmt and c in fmt and isinstance(v, (int, float)):
                cell.number_format = fmt[c]
            if isinstance(v, str):
                f = _status_font(v)
                if f:
                    cell.font = f
    return top + len(rows) + 1


def _title(ws, text: str, note: str = ""):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=13)
    if note:
        ws["A2"] = note


def write_results(path: str, project: Project, result) -> None:
    """解析結果のシートを書き直し、History に1行追記する。"""
    wb = load_workbook(path)
    for name in RESULT_SHEETS:
        if name in wb.sheetnames:
            del wb[name]
    if isinstance(result, ScreeningResult):
        _write_screening(wb, project, result)
        summary = f"スクリーニング解析：有意な因子 {', '.join(result.active) or 'なし'}"
    else:
        _write_rsm(wb, project, result)
        rob = result.optimization.robust
        summary = ("2次モデル解析：" + ("判断ゲートに×あり" if result.blocked else "判断ゲートに×なし") +
                   ("" if rob is None else f"、ロバスト最適 余裕{rob.margin:.3g}"))
    if S_HISTORY in wb.sheetnames:
        wb.move_sheet(S_HISTORY, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(S_HISTORY))
    _append_history(wb, "解析", summary)
    wb.save(path)


def _real_headers(project: Project) -> List[str]:
    return [f"{f.name}[{f.unit}]" if f.unit else f.name for f in project.factors]


def _write_rsm(wb, project: Project, res: RSMResult):
    fit, k = res.fit, len(project.factors)
    p4 = "0.0000"

    ws = wb.create_sheet("Model")
    _title(ws, "2次回帰モデル（コード化単位）", "『使用』を0にしてから再解析すると、その項を除外する（モデル縮小）。切片は常に使用。階層性に注意。")
    rows = [[t["term"], int(t["used"]), t["coef"], t["se"], t["t"], t["p"], t["mark"], t["hierarchy"]]
            for t in fit.coefficient_table()]
    end = _table(ws, 3, 1, ["項", "使用(1/0)", "係数", "標準誤差", "t値", "p値", "有意", "階層性"], rows,
                 {2: p4, 3: p4, 4: "0.000", 5: "0.0000"})
    for r in range(4, end):
        _mark_input(ws.cell(r, 2))
    ws.cell(end + 1, 1, "有意：** p<0.01　* p<0.05　△ p<0.1。p値だけで機械的に削らず、技術的な意味と階層性を優先。")
    stats = [["n（実験数）", fit.n], ["p（項数）", fit.n_params], ["残差自由度", fit.df_resid], ["s（残差標準偏差）", fit.s],
             ["R²", fit.r2], ["調整済みR²", fit.r2_adj], ["PRESS", fit.press], ["予測R²", fit.r2_pred],
             ["t値（工程窓・最適化用）", fit.t_window], ["t値（両側・確認実験用）", fit.t_confirm]]
    _table(ws, 3, 10, ["統計量", "値"], stats, {1: "0.0000"})
    _widths(ws, [10, 9, 11, 11, 9, 10, 6, 24, 2, 24, 12])

    ws = wb.create_sheet("ANOVA")
    _title(ws, "分散分析（ANOVA）と適合度の欠如")
    labels = [("回帰", "regression"), ("残差", "residual"), ("　適合度の欠如", "lack_of_fit"),
              ("　純誤差", "pure_error"), ("合計", "total")]
    rows = [[lab] + [fit.anova[key][c] for c in ("df", "ss", "ms", "f", "p")] for lab, key in labels]
    _table(ws, 3, 1, ["要因", "自由度", "平方和", "平均平方", "F値", "p値"], rows,
           {2: "0.0000", 3: "0.0000", 4: "0.000", 5: "0.0000"})
    _widths(ws, [16, 8, 12, 12, 10, 10])

    ws = wb.create_sheet("Gates")
    _title(ws, "判断ゲート", "×がなく、△の内容を理解・記録してから最適化・工程窓へ進む。")
    _table(ws, 3, 1, ["項目", "判定", "内容"], [[g.name, g.status, g.message] for g in res.gates])
    for r in range(4, 4 + len(res.gates)):
        f = _status_font(ws.cell(r, 2).value + " ")
        if f:
            ws.cell(r, 2).font = f
    _widths(ws, [16, 6, 100])

    c = res.canonical
    ws = wb.create_sheet("Canonical")
    _title(ws, "正準解析（停留点の種類と位置）", "λが全て負→最大点、全て正→最小点、正負混在→鞍点。")
    rows = []
    for i, f in enumerate(project.factors):
        xs = None if c.stationary_point is None else c.stationary_point[i]
        rows.append([f.name, xs, None if xs is None else f.to_real(xs), f.unit])
    end = _table(ws, 3, 1, ["因子", "停留点(コード)", "停留点(実値)", "単位"], rows, {1: "0.000", 2: "0.000"})
    info = [["停留点での予測Y", c.y_at_stationary], ["中心からの距離", c.distance],
            ["実験領域内か", None if c.in_region is None else ("領域内" if c.in_region else "領域外（外挿）")]]
    info += [[f"λ{i + 1}", v] for i, v in enumerate(c.eigenvalues)]
    info += [["停留点の種類", c.kind], ["尾根（リッジ）", "あり：ほぼ平らな方向がある" if c.has_ridge else "なし"],
             ["判断", c.advice]]
    _table(ws, end + 1, 1, ["項目", "値"], info, {1: "0.0000"})
    _widths(ws, [18, 16, 16, 8])

    o = res.optimization
    ws = wb.create_sheet("Optimization")
    _title(ws, "最適化（実験領域内のグリッド探索）",
           "ロバスト最適：規格側の予測限界と規格の差（余裕）が最大の条件。原則こちらを推奨条件にする。")
    rows = []
    for i, f in enumerate(project.factors):
        rows.append([f"{f.name}（コード）", o.robust and o.robust.x[i], o.point and o.point.x[i]])
    for i, f in enumerate(project.factors):
        rows.append([_real_headers(project)[i], o.robust and o.robust.x_real[i], o.point and o.point.x_real[i]])
    for lab, attr in (("予測Y", "pred"), ("判定用下限", "lower"), ("判定用上限", "upper"), ("規格への余裕", "margin")):
        rows.append([lab, o.robust and getattr(o.robust, attr), o.point and getattr(o.point, attr)])
    rows.append(["判定"] + [None if cand is None else ("○ 予測限界でも規格を満たす" if cand.meets_spec else
                                                     "× 予測限界では規格を満たせない") for cand in (o.robust, o.point)])
    end = _table(ws, 4, 1, ["項目", "ロバスト最適（推奨）", "点予測の最適"], rows, {1: "0.000", 2: "0.000"})
    for w in o.warnings:
        ws.cell(end, 1, "△ " + w).font = WARN_FONT
        end += 1
    rows = [[i + 1] + list(cand.x_real) + [cand.pred, cand.lower, cand.upper, cand.margin]
            for i, cand in enumerate(o.robust_top)]
    ws.cell(end + 1, 1, "ロバスト最適の上位候補（近い値に集まっていれば最適点は安定。ばらばらなら平坦で窓を広く取れる可能性）").font = BOLD
    _table(ws, end + 2, 1, ["順位"] + _real_headers(project) + ["予測Y", "判定用下限", "判定用上限", "余裕"], rows,
           {j: "0.000" for j in range(1, k + 5)})
    _widths(ws, [24, 20, 20] + [14] * (k + 2))

    ws = wb.create_sheet("OperatingWindow")
    _title(ws, "推定工程窓（予測限界ベース）",
           "確認実験で確かめるまでは『推定』。予測区間は将来の1回の測定値の区間で、量産の工程能力（許容区間・Cpk）の保証ではない。")
    rows = [[w.factor, w.optimum, w.low if w.low is not None else "該当なし", w.high if w.high is not None else "該当なし",
             w.low_real, w.high_real, f.unit] for w, f in zip(res.windows, project.factors)]
    ws.cell(3, 1, "① 1因子ずつの工程窓（他の因子はロバスト最適点に固定）").font = BOLD
    end = _table(ws, 4, 1, ["因子", "最適(コード)", "窓下限(コード)", "窓上限(コード)", "窓下限(実値)", "窓上限(実値)", "単位"],
                 rows, {1: "0.000", 2: "0.000", 3: "0.000", 4: "0.00", 5: "0.00"})
    ws.cell(end, 1, "注意：①は1因子だけを動かした範囲なので楽観的。複数因子が同時にずれると外れることがある → ②で組合せを確認。").font = WARN_FONT
    cc = res.combination
    if cc is not None:
        ws.cell(end + 2, 1, f"② 提案窓の組合せチェック（各因子の下限・中央・上限の全組合せ {len(cc.margin)}点）").font = BOLD
        info = [["評価点数", len(cc.margin)], ["安全限界・領域外の点", cc.n_outside], ["規格外のおそれがある点", cc.n_fail],
                ["最小の余裕", cc.min_margin]]
        info += [[f"最悪の組合せ：{h}", v] for h, v in zip(_real_headers(project), cc.worst_real)]
        info.append(["判定", cc.verdict])
        end = _table(ws, end + 3, 1, ["項目", "値"], info, {1: "0.000"})
        rows = [list(p) + [m, "○" if (m >= 0 and v) else "×"] for p, m, v in zip(cc.points_real, cc.margin, cc.valid)]
        ws.cell(end + 1, 1, "全組合せの余裕").font = BOLD
        _table(ws, end + 2, 1, _real_headers(project) + ["余裕", "判定"], rows, {j: "0.000" for j in range(k + 1)})
    _widths(ws, [24, 14, 14, 14, 14, 14, 8])

    ws = wb.create_sheet("Residuals")
    _title(ws, "残差診断", "スチューデント化残差の絶対値が3超は要確認（測定ミス・異常の可能性）。消す前に原因を調べる。")
    order = project.design.run_order
    rows = []
    for i in range(fit.n):
        sr = fit.studentized_residuals[i]
        rows.append([i + 1, int(order[i]), fit.y[i], fit.fitted[i], fit.residuals[i], fit.leverage[i], sr,
                     "要確認" if abs(sr) > 3 else "注意" if abs(sr) > 2 else ""])
    _table(ws, 3, 1, ["標準順", "実験順", "Y実測", "予測", "残差", "てこ比h", "スチューデント化残差", "判定"], rows,
           {2: "0.000", 3: "0.000", 4: "0.000", 5: "0.000", 6: "0.000"})
    _widths(ws, [8, 8, 10, 10, 10, 10, 18, 8])


def _write_screening(wb, project: Project, res: ScreeningResult):
    fit = res.fit
    ws = wb.create_sheet("Effects")
    _title(ws, "スクリーニング：1次モデルの効果",
           "効果＝係数×2（-1→+1の変化量）。PB計画では交互作用が主効果と交絡するので、絞り込みは確定判断にしない。")
    rows = [[r.term, r.coef, None if r.term == "切片" else r.effect, r.se, r.t, r.p,
             None if r.lenth_active is None else ("有効" if r.lenth_active else "")] for r in fit.effects]
    end = _table(ws, 3, 1, ["項", "係数", "効果", "標準誤差", "t値", "p値", "Lenth法"], rows,
                 {1: "0.0000", 2: "0.0000", 3: "0.0000", 4: "0.000", 5: "0.0000"})
    info = [["残差自由度", fit.df_resid], ["Lenth PSE", fit.lenth_pse], ["Lenth ME（95%）", fit.lenth_me],
            ["有意な因子", "、".join(res.active) or "なし"]]
    cv = fit.curvature
    if cv is not None:
        info += [["要因点の平均", cv.mean_factorial], ["中心点の平均", cv.mean_center], ["曲率の平方和", cv.ss],
                 ["曲率のF値", cv.f], ["曲率のp値", cv.p], ["曲率の判定", cv.message]]
    else:
        info.append(["曲率の判定", "中心点がないため検定不可"])
    info += [["注記", n] for n in fit.notes]
    _table(ws, end + 1, 1, ["項目", "値"], info, {1: "0.0000"})
    _widths(ws, [16, 14, 12, 12, 10, 10, 10])

    if res.path:
        ws = wb.create_sheet("SteepestPath")
        goal = "最急上昇" if project.response.goal != "minimize" else "最急降下"
        _title(ws, f"{goal}の経路",
               "経路に沿って実験し、応答が下がり始めた手前を新しい中心にして次の計画（2水準＋中心点、またはCCD）を組む。")
        k = len(project.factors)
        rows = [[s.step] + list(s.x) + list(s.x_real) + ["× 安全限界超過" if s.outside_safety else "", None]
                for s in res.path]
        _table(ws, 3, 1, ["ステップ"] + [f"X{i + 1}(コード)" for i in range(k)] + _real_headers(project) +
               ["安全限界", "Y実測（記録用）"], rows, {j: "0.000" for j in range(1, 2 * k + 1)})
        for r in range(4, 4 + len(rows)):
            _mark_input(ws.cell(r, 2 * k + 3))
        _widths(ws, [9] + [11] * k + [14] * k + [16, 16])


# ---------------------------------------------------------------- 既存Excel（RSMツール v1.1）の取り込み

def import_excel_v1(path: str) -> Project:
    """RSMツール（Excel）v1.1 のブックから、因子・目的変数・計画・Yを読み込む。"""
    wb = load_workbook(path, data_only=True)
    s1 = wb["01_因子設定"]
    factors = []
    r = 5
    while isinstance(s1.cell(r, 1).value, (int, float)):
        factors.append(Factor(s1.cell(r, 2).value, _num(s1.cell(r, 5).value), _num(s1.cell(r, 6).value),
                              s1.cell(r, 4).value or "", _num(s1.cell(r, 9).value), _num(s1.cell(r, 10).value),
                              s1.cell(r, 11).value == "○"))
        r += 1
    setting = {}
    for r in range(10, 30):
        if s1.cell(r, 1).value:
            setting[s1.cell(r, 1).value] = s1.cell(r, 2).value
    goal = _inv(GOAL_LABELS)[setting["最適化の方向"]]
    response = Response(setting.get("名称") or "Y", goal, _num(setting.get("下限規格 LSL")),
                        _num(setting.get("上限規格 USL")), _num(setting.get("目標値 T")),
                        _num(setting.get("予測区間の信頼水準")) or 0.95, setting.get("単位") or "")
    label = str(setting["計画の種類"])
    design_type = ("face_ccd" if label.startswith("面心") else "rotatable_ccd" if label.startswith("回転可能")
                   else "box_behnken")
    randomization = "grouped" if setting.get("ランダマイズ方式") == "変更困難因子でグループ化" else "complete"
    project = Project(f"{path.split('/')[-1]} から取り込み", factors, response, design_type,
                      int(setting["中心点の数"]), 1, randomization)
    d = project.make_design()

    s2 = wb["02_実験計画"]
    hdr = {s2.cell(4, c).value: c for c in range(1, s2.max_column + 1) if s2.cell(4, c).value}
    y_col = next(c for h, c in hdr.items() if str(h).startswith("Y"))
    order = np.zeros(d.n_runs, int)
    y = np.full(d.n_runs, np.nan)
    k = len(factors)
    for r in range(5, s2.max_row + 1):
        if s2.cell(r, 1).value is None:
            break
        if s2.cell(r, 3).value != 1:
            continue
        i = int(s2.cell(r, 1).value) - 1
        x = [float(s2.cell(r, 4 + j).value) for j in range(k)]
        if not np.allclose(x, d.points[i]):
            raise ValueError(f"標準順{i + 1}のコード値がv1.1の計画と合わない")
        order[i] = int(s2.cell(r, hdr["実験順序"]).value)
        v = _num(s2.cell(r, y_col).value)
        y[i] = np.nan if v is None else v
    d.run_order = order
    d.seed = None
    project.y = y
    s4 = wb["04_回帰ANOVA"]
    names = term_names(k)
    used = [int(s4.cell(6 + i, 2).value) for i in range(len(names))]
    project.used_terms = used
    return project


def save_project(path: str, project: Project) -> None:
    """Project（取り込んだものなど）を新しいプロジェクトファイルに書き出す。計画とYも含める。"""
    write_template(path, project=project)
    wb = load_workbook(path)
    _write_design_sheet(wb, project, f"実験指示書（{DESIGN_LABELS[project.design_type]}・{project.design.n_runs}回）")
    _append_history(wb, "取り込み", "既存の計画とYを書き出し")
    wb.save(path)
