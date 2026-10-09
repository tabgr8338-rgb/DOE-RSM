"""画面で操作する DOE-RSM（Streamlit）。起動：python -m doe_rsm app

1つの実験テーマを1つのプロジェクトファイル（Excel）で扱う。
設定 → 実験計画・Y入力 → 解析 → 最適条件・工程窓（スクリーニングなら次の実験へ）→ 確認実験 の順に進む。
計算は既存のエンジン（Excel v1.1 と一致を確認済み）を使い、画面は入力と表示だけを受け持つ。
"""
import tempfile
from dataclasses import replace
from pathlib import Path
from typing import List, Optional

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from . import __version__
from .charts import contour, pareto, residual_plots
from .confirm import judge_confirmation
from .design import BOX_BEHNKEN, FACE_CCD
from .excel_io import export_project, import_excel_v1, read_project
from .factors import Factor, Response
from .model import term_names
from .project import (DESIGN_LABELS, GOAL_LABELS, RANDOMIZATION_LABELS, RSM_TYPES, Project, RSMResult,
                      ScreeningResult, analyze)
from .screening import PLACKETT_BURMAN, pb_runs_for, recenter

SAMPLE = Path(__file__).parent / "samples" / "sample_3factor.xlsx"
GOAL_KEYS = list(GOAL_LABELS)
DESIGN_KEYS = list(DESIGN_LABELS)
RAND_KEYS = list(RANDOMIZATION_LABELS)
FACTOR_COLS = ["因子名", "単位", "下限(-1)", "上限(+1)", "安全下限", "安全上限", "変更困難"]
STATUS_ICON = {"○": "✅", "△": "⚠️", "×": "❌", "－": "➖"}


# ---------------------------------------------------------------- 状態とファイル

def _state():
    ss = st.session_state
    ss.setdefault("project", None)
    ss.setdefault("v", 0)              # 読み込み・新規作成のたびに増やし、入力欄を作り直す
    ss.setdefault("file_name", None)
    ss.setdefault("n_new_factors", 3)
    ss.setdefault("flash", None)
    return ss


def _set_project(project: Optional[Project], file_name: Optional[str], flash: Optional[str] = None):
    ss = st.session_state
    ss.project = project
    ss.file_name = file_name
    ss.v += 1
    ss.flash = flash


def load_file(name: str, data: bytes) -> Project:
    """本ツールのプロジェクトファイル、または RSMツール v1.1 のブックを読む。"""
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / name
        path.write_bytes(data)
        wb = load_workbook(path, read_only=True)
        sheets = wb.sheetnames
        wb.close()  # 読み取り専用モードはファイルを開いたままにするので、Windows で一時フォルダを消せるよう閉じる
        if "01_因子設定" in sheets:
            project = import_excel_v1(str(path))
            project.log("取り込み", f"RSMツール v1.1「{name}」から計画とYを取り込み")
            return project
        if "Project" not in sheets or "Factors" not in sheets:
            raise ValueError("DOE-RSM のプロジェクトファイルでも、RSMツール v1.1 のブックでもありません")
        return read_project(str(path))


def project_bytes(project: Project, result=None) -> bytes:
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "project.xlsx"
        export_project(str(path), project, result)
        return path.read_bytes()


def _design_signature(p: Project):
    """変えると実験計画を作り直す必要がある設定。"""
    fac = tuple((f.name, f.low, f.high, f.hard_to_change if p.randomization == "grouped" else None)
                for f in p.factors)
    return fac, p.design_type, p.n_center, p.seed, p.randomization, p.pb_runs if p.is_screening else None


def _num(v) -> Optional[float]:
    if v is None or (isinstance(v, float) and np.isnan(v)) or v == "":
        return None
    return float(v)


def _fmt(v, digits: int = 4) -> str:
    return "" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:.{digits}g}"


def _real_cond(factors: List[Factor], x_real) -> str:
    return "、".join(f"{f.name} {v:.4g}{f.unit}" for f, v in zip(factors, x_real))


# ---------------------------------------------------------------- サイドバー

def sidebar(ss):
    st.sidebar.title("DOE-RSM")
    st.sidebar.caption(f"実験計画・応答曲面法　v{__version__}")
    up = st.sidebar.file_uploader("プロジェクトファイルを開く（.xlsx）", type=["xlsx"],
                                  help="このツールで保存したファイル、またはRSMツール v1.1 のブック")
    if up is not None and ss.get("uploaded_id") != (up.name, up.size):
        ss.uploaded_id = (up.name, up.size)
        try:
            _set_project(load_file(up.name, up.getvalue()), up.name, f"「{up.name}」を開きました")
        except Exception as e:  # noqa: BLE001 - 読めない理由をそのまま利用者に見せる
            st.sidebar.error(f"開けませんでした：{e}")
    c1, c2 = st.sidebar.columns([1, 1])
    n = c1.number_input("因子数", 2, 23, ss.n_new_factors, key="n_new")
    if c2.button("新規作成", width="stretch"):
        ss.n_new_factors = int(n)
        _set_project(None, None, "新しいプロジェクトです。①設定 を入力してください")
    if SAMPLE.exists() and st.sidebar.button("サンプルで試す（3因子・面心CCD）", width="stretch"):
        _set_project(read_project(str(SAMPLE)), "sample_3factor.xlsx", "サンプル（架空データ）を開きました")



def sidebar_save(ss):
    """画面の入力（Yなど）をすべて反映したあとで呼ぶ。"""
    p = ss.project
    if p is None:
        return
    result = None
    if p.design is not None and p.y is not None and p.n_missing == 0:
        try:
            result = analyze(p)
        except ValueError:  # 解析できない状態でも設定と計画は保存できるようにする
            result = None
    name = ss.file_name or f"{p.title or 'project'}.xlsx"
    st.sidebar.download_button("💾 プロジェクトファイルを保存", project_bytes(p, result), file_name=name,
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               width="stretch", type="primary",
                               help="設定・実験計画・Y・解析結果・履歴を1つのExcelに保存します")
    st.sidebar.caption("保存したファイルは次回「開く」で続きから使えます。Excelで開いてYを入力しても構いません。")


# ---------------------------------------------------------------- ① 設定

def tab_settings(ss):
    p: Optional[Project] = ss.project
    v = ss.v
    st.subheader("目的変数と計画")
    c1, c2, c3 = st.columns([2, 2, 1])
    title = c1.text_input("テーマ", p.title if p else "", key=f"title{v}")
    rname = c2.text_input("目的変数", p.response.name if p else "", key=f"rname{v}", placeholder="例：剥離強度")
    runit = c3.text_input("単位", p.response.unit if p else "", key=f"runit{v}")

    c1, c2, c3, c4, c5 = st.columns(5)
    goal = c1.selectbox("最適化の方向", GOAL_KEYS, GOAL_KEYS.index(p.response.goal) if p else 0,
                        format_func=GOAL_LABELS.get, key=f"goal{v}")
    lsl = c2.number_input("下限規格 LSL", value=p.response.lsl if p else None, format="%g", key=f"lsl{v}",
                          disabled=goal == "minimize")
    usl = c3.number_input("上限規格 USL", value=p.response.usl if p else None, format="%g", key=f"usl{v}",
                          disabled=goal == "maximize")
    target = c4.number_input("目標値 T", value=p.response.target if p else None, format="%g", key=f"t{v}",
                             disabled=goal != "target")
    conf = c5.number_input("予測区間の信頼水準", 0.5, 0.999, p.response.confidence if p else 0.95, 0.01,
                           format="%.3g", key=f"conf{v}",
                           help="最大化・最小化は片側、目標値は両側で判定。確認実験は常に両側")

    c1, c2, c3, c4, c5 = st.columns(5)
    dtype = c1.selectbox("計画の種類", DESIGN_KEYS, DESIGN_KEYS.index(p.design_type) if p else 0,
                         format_func=DESIGN_LABELS.get, key=f"dtype{v}",
                         help="スクリーニング：Plackett-Burman／2水準要因　応答曲面：面心CCD／回転可能CCD／Box-Behnken")
    n_center = c2.number_input("中心点の数", 0, 20, p.n_center if p else 6, key=f"nc{v}",
                               help="純誤差・曲率の評価に使うので3以上。RSMでは5〜6を推奨")
    pb_opts = [0, 8, 12, 16, 20, 24]
    pb_runs = c3.selectbox("PB実験数", pb_opts, pb_opts.index(p.pb_runs) if p and p.pb_runs in pb_opts else 0,
                           format_func=lambda x: "自動" if x == 0 else f"{x}回", key=f"pb{v}",
                           disabled=dtype != PLACKETT_BURMAN) or None
    rand = c4.selectbox("ランダマイズ", RAND_KEYS, RAND_KEYS.index(p.randomization) if p else 0,
                        format_func=RANDOMIZATION_LABELS.get, key=f"rand{v}")
    seed = c5.number_input("乱数シード", 0, 10 ** 9, p.seed if p else 1, key=f"seed{v}",
                           help="同じシードなら同じ実験順序を再現できる")

    st.subheader("因子")
    st.caption("下限・上限はコード値±1に当たる実値。安全限界は空欄なら制限なし。行の追加・削除もできます。")
    if p:
        rows = [[f.name, f.unit, f.low, f.high, f.safety_low, f.safety_high, f.hard_to_change] for f in p.factors]
    else:
        rows = [[f"因子{i + 1}", "", None, None, None, None, False] for i in range(ss.n_new_factors)]
    df = pd.DataFrame(rows, columns=FACTOR_COLS)
    num_cols = ["下限(-1)", "上限(+1)", "安全下限", "安全上限"]
    df[num_cols] = df[num_cols].astype(float)   # 空欄だけの列が object 型になると入力が反映されない
    df["単位"] = df["単位"].fillna("").astype(str)
    df["変更困難"] = df["変更困難"].astype(bool)
    edited = st.data_editor(df, num_rows="dynamic", width="stretch", key=f"factors{v}",
                            column_config={
                                "下限(-1)": st.column_config.NumberColumn(format="%g"),
                                "上限(+1)": st.column_config.NumberColumn(format="%g"),
                                "安全下限": st.column_config.NumberColumn(format="%g"),
                                "安全上限": st.column_config.NumberColumn(format="%g"),
                                "変更困難": st.column_config.CheckboxColumn(help="段取り替えに時間がかかる因子"),
                            })

    try:
        factors = []
        for _, r in edited.iterrows():
            name = r["因子名"]
            if name is None or str(name).strip() == "" or (isinstance(name, float) and np.isnan(name)):
                continue
            low, high = _num(r["下限(-1)"]), _num(r["上限(+1)"])
            if low is None or high is None or high <= low:
                raise ValueError(f"因子「{name}」の下限・上限を入力してください（下限 < 上限）")
            factors.append(Factor(str(name), low, high, "" if pd.isna(r["単位"]) else str(r["単位"]),
                                  _num(r["安全下限"]), _num(r["安全上限"]), bool(r["変更困難"])))
        if len(factors) < 2:
            raise ValueError("因子を2つ以上入力してください")
        k = len(factors)
        if dtype in RSM_TYPES and not 2 <= k <= 4:
            raise ValueError("応答曲面（CCD・Box-Behnken）は2〜4因子です。多いときはまずスクリーニングで絞り込みます")
        if dtype == BOX_BEHNKEN and k < 3:
            raise ValueError("Box-Behnken は3因子以上です")
        if dtype == PLACKETT_BURMAN and pb_runs is not None and pb_runs - 1 < k:
            raise ValueError(f"{pb_runs}回のPB計画には{pb_runs - 1}因子まで（自動なら{pb_runs_for(k)}回）")
        response = Response(rname or "Y", goal, lsl if goal != "minimize" else None,
                            usl if goal != "maximize" else None, target if goal == "target" else None, conf, runit)
        new = Project(title, factors, response, dtype, int(n_center), int(seed), rand, pb_runs)
    except ValueError as e:
        st.warning(str(e))
        return

    if p is not None:
        new.history = p.history
    same_design = p is not None and p.design is not None and _design_signature(p) == _design_signature(new)
    has_y = p is not None and p.y is not None and p.n_missing < len(p.y)
    if same_design:
        label = "設定を反映"
        st.caption("実験計画に関わる設定は変わっていません。規格や安全限界などの変更だけを反映し、入力済みのYは残します。")
    elif p is not None and p.design is not None and has_y:
        st.warning("因子・計画の種類・中心点・シードなどを変えたので、実験計画を作り直します。入力済みのYは消えます。")
        if not st.checkbox("入力済みのYを消して作り直す", key=f"confirm_reset{v}"):
            return
        label = "実験計画を作り直す"
    else:
        label = "実験計画を作る"
    if st.button(label, type="primary"):
        if same_design:
            new.design, new.y, new.used_terms = p.design, p.y, p.used_terms
            new.design.factors = new.factors
            new.log("設定", "規格・安全限界などの設定を変更")
            ss.project = new
            ss.v += 1
            ss.flash = "設定を反映しました"
        else:
            if p is None:
                new.log("作成", "プロジェクトを作成")
            d = new.make_design()
            for msg in d.notes:
                new.log("注意", msg)
            new.log("計画", f"{DESIGN_LABELS[dtype]} {d.n_runs}回、中心点{n_center}、{RANDOMIZATION_LABELS[rand]}、シード{seed}")
            ss.project = new
            ss.v += 1
            ss.flash = f"{d.n_runs}回の実験計画を作りました。②実験計画・Y入力 に進んでください"
        st.rerun()


# ---------------------------------------------------------------- ② 実験計画・Y入力

def tab_design(ss):
    p: Optional[Project] = ss.project
    if p is None or p.design is None:
        st.info("①設定 で「実験計画を作る」を押すと、ここに実験指示書が出ます。")
        return
    d = p.design
    st.markdown(f"**{DESIGN_LABELS[p.design_type]}・{d.n_runs}回**　ランダマイズ：{RANDOMIZATION_LABELS[p.randomization]}"
                + (f"・乱数シード {d.seed}" if d.seed is not None else ""))
    if d.is_safe:
        st.success("全実験点が安全限界内です")
    else:
        st.error("安全限界を超える実験点があります。この計画は実行できません。範囲を縮小するか計画の種類を変えてください。")
        st.dataframe(pd.DataFrame([{"因子": r["factor"], "計画の最小": r["min"], "計画の最大": r["max"],
                                    "安全下限": r["safety_low"], "安全上限": r["safety_high"],
                                    "判定": "○" if r["ok"] else "×"} for r in d.safety_check()]), hide_index=True)
    for n in d.notes:
        st.warning(n)

    st.caption("上から順に実験し、Y の列に結果を入力してください（Excelからまとめて貼り付けもできます）。")
    order = np.argsort(d.run_order, kind="stable")
    real = d.real_points()
    y = p.y if p.y is not None else np.full(d.n_runs, np.nan)
    ycol = f"Y：{p.response.name}" + (f"[{p.response.unit}]" if p.response.unit else "")
    data = {"実験順": [int(d.run_order[i]) for i in order], "標準順": [int(i + 1) for i in order],
            "区分": [d.kinds[i] for i in order]}
    for j, f in enumerate(p.factors):
        data[f"{f.name}" + (f"[{f.unit}]" if f.unit else "")] = [float(real[i, j]) for i in order]
    data[ycol] = [float(y[i]) for i in order]
    df = pd.DataFrame(data)
    df[ycol] = df[ycol].astype(float)
    edited = st.data_editor(df, hide_index=True, width="stretch", key=f"y{ss.v}",
                            disabled=[c for c in df.columns if c != ycol],
                            column_config={ycol: st.column_config.NumberColumn(format="%g")},
                            height=min(38 * (len(df) + 1) + 3, 800))
    new_y = np.full(d.n_runs, np.nan)
    for std, val in zip(edited["標準順"], edited[ycol]):
        v = _num(val)
        new_y[int(std) - 1] = np.nan if v is None else v
    p.y = new_y
    if p.n_missing:
        st.info(f"Y が {p.n_missing} 件未入力です。すべて入力すると ③解析 が出ます。")
    else:
        st.success("全件入力済みです。③解析 に進んでください。")


# ---------------------------------------------------------------- ③ 解析

def _model_terms(ss, p: Project):
    names = term_names(len(p.factors))
    used = p.used_terms or [1] * len(names)
    with st.expander("モデルの項を選ぶ（モデル縮小）", expanded=False):
        st.caption("p値が大きい項を外すとモデルが単純になります。主効果を外して交互作用・2乗項を残すと階層性の注意が出ます。切片は常に使います。")
        df = pd.DataFrame({"項": names, "使用": [bool(u) for u in used]})
        ed = st.data_editor(df, hide_index=True, disabled=["項"], key=f"terms{ss.v}", width="content")
        new = [1] + [int(b) for b in ed["使用"].tolist()[1:]]
        if new != list(used):
            p.used_terms = None if all(new) else new
            p.log("モデル", "使用する項を変更：" + "、".join(n for n, u in zip(names, new) if not u) + " を除外"
                  if not all(new) else "全項を使用")


def _result(ss):
    p: Optional[Project] = ss.project
    if p is None or p.design is None:
        return None, "①設定 で実験計画を作ってください。"
    if p.y is None or p.n_missing:
        return None, f"Y が {p.n_missing if p.y is not None else p.design.n_runs} 件未入力です。②実験計画・Y入力 で入力してください。"
    try:
        return analyze(p), None
    except ValueError as e:
        return None, f"解析できません：{e}"


def tab_analysis(ss):
    p: Optional[Project] = ss.project
    if p is not None and p.design is not None and not p.is_screening:
        _model_terms(ss, p)
    res, why = _result(ss)
    if res is None:
        st.info(why)
        return
    if isinstance(res, ScreeningResult):
        _screening_analysis(p, res)
    else:
        _rsm_analysis(p, res)


def _rsm_analysis(p: Project, res: RSMResult):
    fit = res.fit
    st.subheader("判断ゲート")
    if res.blocked:
        st.error("× があるため、最適化には進まないでください（④は参考表示になります）。")
    st.markdown("\n".join(f"- {STATUS_ICON.get(g.status, '')} **{g.name}**：{g.message}" for g in res.gates))

    c = st.columns(5)
    c[0].metric("R²", f"{fit.r2:.3f}")
    c[1].metric("調整済み R²", f"{fit.r2_adj:.3f}")
    c[2].metric("予測 R²", f"{fit.r2_pred:.3f}")
    c[3].metric("残差の標準偏差 s", _fmt(fit.s))
    c[4].metric("PRESS", _fmt(fit.press))

    st.subheader("回帰係数（コード化単位）")
    st.dataframe(pd.DataFrame([{"項": r["term"], "使用": "○" if r["used"] else "", "係数": r["coef"] if r["used"] else None,
                                "標準誤差": r["se"], "t値": r["t"], "p値": r["p"], "": r["mark"], "階層性": r["hierarchy"]}
                               for r in fit.coefficient_table()]),
                 hide_index=True, width="stretch",
                 column_config={k: st.column_config.NumberColumn(format="%.4g") for k in ("係数", "標準誤差", "t値", "p値")})
    st.caption("　".join(f"X{i + 1}＝{f.name}" for i, f in enumerate(p.factors)) + "　｜　** p<0.01　* p<0.05　△ p<0.10")

    st.subheader("分散分析（ANOVA）")
    labels = {"regression": "回帰", "residual": "残差", "lack_of_fit": "　適合度の欠如", "pure_error": "　純誤差",
              "total": "全体"}
    st.dataframe(pd.DataFrame([{"要因": labels[k], "自由度": a["df"], "平方和": a["ss"], "平均平方": a["ms"], "F値": a["f"],
                                "p値": a["p"]} for k, a in fit.anova.items()]),
                 hide_index=True, width="stretch",
                 column_config={k: st.column_config.NumberColumn(format="%.4g") for k in ("平方和", "平均平方", "F値", "p値")})

    st.subheader("残差")
    st.plotly_chart(residual_plots(fit), width="stretch")

    st.subheader("正準解析（停留点）")
    can = res.canonical
    st.write(f"曲面の形：**{can.kind}**" + ("（リッジあり：最適条件が一点に決まらず、尾根・谷に沿って同程度の条件が続く）"
                                          if can.has_ridge else ""))
    if can.stationary_point is not None:
        sp_real = [f.to_real(v) for f, v in zip(p.factors, can.stationary_point)]
        st.write(f"停留点：{_real_cond(p.factors, sp_real)}　予測 {can.y_at_stationary:.4g}　"
                 f"中心からの距離 {can.distance:.3g}（{'領域内' if can.in_region else '領域外'}）")
    st.write("固有値：" + "、".join(f"{e:.4g}" for e in can.eigenvalues))
    st.info(can.advice)


def _screening_analysis(p: Project, res: ScreeningResult):
    fit = res.fit
    st.subheader("有意な因子")
    st.write("、".join(res.active) if res.active else "有意な因子は見つかりませんでした（範囲を広げる・因子を見直す）")
    for n in fit.notes:
        st.caption(n)
    st.plotly_chart(pareto(fit, p.factors), width="stretch")
    names = {f"X{i + 1}": f.name for i, f in enumerate(p.factors)}
    st.dataframe(pd.DataFrame([{"因子": names.get(e.term, e.term), "係数": e.coef,
                                "効果": None if e.term == "切片" else e.effect, "標準誤差": e.se, "t値": e.t,
                                "p値": e.p, "Lenth判定": "" if e.lenth_active is None else ("有意" if e.lenth_active else "")}
                               for e in fit.effects]), hide_index=True, width="stretch",
                 column_config={k: st.column_config.NumberColumn(format="%.4g")
                                for k in ("係数", "効果", "標準誤差", "t値", "p値")})
    st.subheader("曲率（中心点）")
    cv = fit.curvature
    if cv is None:
        st.info("中心点がないため曲率を判定できません")
    else:
        st.write(f"要因点の平均 {cv.mean_factorial:.4g}、中心点の平均 {cv.mean_center:.4g}" +
                 ("" if cv.p is None else f"、p値 {cv.p:.3g}"))
        (st.error if cv.significant else st.success if cv.significant is False else st.info)(cv.message)


# ---------------------------------------------------------------- ④ 最適条件・工程窓 / 次の実験へ

def tab_optimum(ss):
    res, why = _result(ss)
    if res is None:
        st.info(why)
        return
    p: Project = ss.project
    if isinstance(res, ScreeningResult):
        _next_experiment(ss, p, res)
        return
    if res.blocked:
        st.error("判断ゲートに × があります。以下は参考値です。原因を解消してから使ってください。")
    opt = res.optimization
    for w in opt.warnings:
        st.warning(w)
    if opt.robust is None:
        st.error("領域内に規格を満たす条件が見つかりません。")
        return
    r, pt = opt.robust, opt.point
    unit = p.response.unit
    st.subheader("推奨条件")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**ロバスト最適**（予測限界でも規格から最も余裕がある条件）")
        st.write(_real_cond(p.factors, r.x_real))
        st.write(f"予測 {r.pred:.4g}{unit}　予測限界 {r.lower:.4g}〜{r.upper:.4g}　規格からの余裕 {r.margin:.3g}")
    with c2:
        st.markdown("**点最適**（予測値そのものが最も良い条件）")
        st.write(_real_cond(p.factors, pt.x_real))
        st.write(f"予測 {pt.pred:.4g}{unit}　予測限界 {pt.lower:.4g}〜{pt.upper:.4g}　余裕 {pt.margin:.3g}")
    with st.expander("ロバスト最適の上位候補"):
        st.dataframe(pd.DataFrame([{**{f.name: v for f, v in zip(p.factors, c.x_real)}, "予測": c.pred,
                                    "予測下限": c.lower, "予測上限": c.upper, "余裕": c.margin} for c in opt.robust_top]),
                     hide_index=True, width="stretch")

    st.subheader("等高線（2因子の断面）")
    k = len(p.factors)
    names = [f.name for f in p.factors]
    c1, c2 = st.columns(2)
    xi = c1.selectbox("横軸", range(k), 0, format_func=names.__getitem__, key=f"cx{ss.v}")
    yi = c2.selectbox("縦軸", [i for i in range(k) if i != xi], 0, format_func=names.__getitem__, key=f"cy{ss.v}_{xi}")
    fixed = np.array(r.x, float)
    others = [i for i in range(k) if i not in (xi, yi)]
    if others:
        cols = st.columns(len(others))
        for c, i in zip(cols, others):
            f = p.factors[i]
            val = c.slider(f"{f.name} を固定", float(f.to_real(-1)), float(f.to_real(1)), float(r.x_real[i]),
                           format="%g", key=f"fix{ss.v}_{i}")
            fixed[i] = f.to_coded(val)
    st.plotly_chart(contour(res.fit, p.response, p.factors, p.design_type, xi, yi, fixed, p.design.points, r.x),
                    width="stretch")
    st.caption("赤線は「予測限界でちょうど規格」の境界です。★はロバスト最適。白丸はこの断面上の実験点。")

    st.subheader("工程窓（推定）")
    st.caption("ロバスト最適から1因子だけを動かして規格を満たす範囲。最後に全組合せ（3^k点）で確かめます。確認実験の前は「推定」です。")
    st.dataframe(pd.DataFrame([{"因子": w.factor, "最適": f.to_real(w.optimum),
                                "下限": w.low_real, "上限": w.high_real,
                                "": "" if w.low is not None else "該当なし（中心に固定）"}
                               for w, f in zip(res.windows, p.factors)]), hide_index=True, width="stretch")
    cc = res.combination
    if cc is not None:
        (st.success if cc.ok else st.error)("組合せチェック：" + cc.verdict)
        if not cc.ok:
            st.write(f"最も厳しい組合せ：{_real_cond(p.factors, cc.worst_real)}（余裕 {cc.min_margin:.3g}）")


def _next_experiment(ss, p: Project, res: ScreeningResult):
    fit = res.fit
    cv = fit.curvature
    if cv is not None and cv.significant:
        st.warning("曲率が出ています。この領域で応答曲面（CCD・Box-Behnken）を組むのが次の一手です。")
    if not res.path:
        st.info("有意な因子がないため、最急上昇の経路は出せません。")
        return
    goal = "最急上昇" if p.response.goal == "maximize" else "最急降下"
    st.subheader(f"{goal}の経路")
    st.caption("係数が最大の因子をコード値で1ずつ動かし、他の因子は係数の比で動かします。"
               "経路に沿って実験し、応答が下がり始めた手前を中心に次の計画を組みます。")
    rows = [{"ステップ": s.step, **{f.name: v for f, v in zip(p.factors, s.x_real)},
             "安全限界": "× 超過" if s.outside_safety else ""} for s in res.path]
    st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch")

    st.subheader("次の計画（応答曲面）を作る")
    c1, c2, c3 = st.columns(3)
    step = c1.selectbox("中心にするステップ", [s.step for s in res.path], 0, key=f"step{ss.v}")
    names = [f.name for f in p.factors]
    keep = c2.multiselect("残す因子（2〜4）", names, res.active[:4] or names[:2], key=f"keep{ss.v}")
    dtype = c3.selectbox("計画の種類", list(RSM_TYPES), 0, format_func=DESIGN_LABELS.get, key=f"ndt{ss.v}")
    if not 2 <= len(keep) <= 4:
        st.info("残す因子を2〜4個選んでください。")
        return
    if dtype == BOX_BEHNKEN and len(keep) < 3:
        st.info("Box-Behnken は3因子以上です。")
        return
    s = res.path[step]
    idx = [names.index(n) for n in keep]
    fixed = [f"{p.factors[i].name} {s.x_real[i]:.4g}{p.factors[i].unit}" for i in range(len(names)) if i not in idx]
    if fixed:
        st.caption("外す因子は次の値に固定して実験します：" + "、".join(fixed))
    if st.button("この条件で新しいプロジェクトを作る", type="primary"):
        factors = recenter([p.factors[i] for i in idx], [s.x_real[i] for i in idx])
        new = Project(f"{p.title}（応答曲面）" if p.title else "応答曲面", factors, replace(p.response), dtype, 6, p.seed,
                      "complete")
        new.log("作成", f"「{p.title}」のスクリーニングから。{goal}ステップ{step}を中心に設定"
                + ("。固定：" + "、".join(fixed) if fixed else ""))
        d = new.make_design()
        new.log("計画", f"{DESIGN_LABELS[dtype]} {d.n_runs}回、中心点6、完全ランダム、シード{p.seed}")
        _set_project(new, None, "応答曲面の新しいプロジェクトを作りました。元のスクリーニングは保存しておいてください。"
                                "範囲は①設定で調整できます")
        st.rerun()


# ---------------------------------------------------------------- ⑤ 確認実験

def tab_confirm(ss):
    res, why = _result(ss)
    if res is None:
        st.info(why)
        return
    if isinstance(res, ScreeningResult):
        st.info("確認実験は応答曲面で推奨条件を出したあとに行います。")
        return
    p: Project = ss.project
    fit = res.fit
    st.caption("推奨条件で同じ条件の実験を3回以上行い、実測値が予測区間（両側）に入るかを確かめます。")
    base = res.optimization.robust.x_real if res.optimization.robust is not None else [f.center for f in p.factors]
    cols = st.columns(len(p.factors))
    x_real = [c.number_input(f"{f.name}" + (f"[{f.unit}]" if f.unit else ""), value=float(v), format="%g",
                             key=f"cf{ss.v}_{i}") for i, (c, f, v) in enumerate(zip(cols, p.factors, base))]
    obs_df = st.data_editor(pd.DataFrame({"実測値": [None, None, None]}, dtype=float), num_rows="dynamic",
                            key=f"obs{ss.v}", column_config={"実測値": st.column_config.NumberColumn(format="%g")})
    obs = [float(v) for v in obs_df["実測値"] if _num(v) is not None]
    coded = np.array([f.to_coded(v) for f, v in zip(p.factors, x_real)])
    if np.abs(coded).max() > 1 + 1e-9:
        st.warning("実験した範囲（±1）の外の条件です。予測は外挿になります。")
    runs = judge_confirmation(fit, p.factors, x_real, obs or [np.nan])
    r0 = runs[0]
    st.write(f"予測 {r0.pred:.4g}　予測区間（{fit.confidence:.0%}・両側） {r0.pi_low:.4g}〜{r0.pi_high:.4g}")
    if not obs:
        return
    st.dataframe(pd.DataFrame([{"実測値": r.observed, "予測区間内": "○" if r.model_ok else "×",
                                "規格": "○" if r.spec_ok(p.response) else "×"} for r in runs]), hide_index=True)
    mean = float(np.mean(obs))
    all_in = all(r.model_ok for r in runs)
    all_spec = all(r.spec_ok(p.response) for r in runs)
    if all_in and all_spec:
        verdict = "○ 実測が予測区間内で規格も満たす → モデルを採用し、推奨条件と工程窓を確定してよい"
        st.success(verdict)
    elif all_in:
        verdict = "△ モデルどおりだが規格外の実測がある → 余裕の大きい条件を選ぶか、ばらつきを下げる"
        st.warning(verdict)
    else:
        verdict = "× 予測区間から外れた → モデルが現場を表していない。条件の再現性・測定・計画を見直す"
        st.error(verdict)
    st.write(f"平均 {mean:.4g}（{len(obs)}回）")
    if len(obs) < 3:
        st.caption("確認実験は3回以上を推奨します。")
    if st.button("この結果を履歴に記録"):
        p.log("確認実験", f"{_real_cond(p.factors, x_real)}：実測 {', '.join(f'{v:.4g}' for v in obs)}"
                          f"（予測区間 {r0.pi_low:.4g}〜{r0.pi_high:.4g}）{verdict}")
        st.toast("履歴に記録しました。保存すると History シートに残ります。")


# ---------------------------------------------------------------- 画面全体

def main():
    st.set_page_config(page_title="DOE-RSM", page_icon="📈", layout="wide")
    ss = _state()
    sidebar(ss)
    p = ss.project
    st.title(p.title if p and p.title else "DOE-RSM")
    if ss.flash:
        st.success(ss.flash)
        ss.flash = None
    screening = p is not None and p.is_screening
    tabs = st.tabs(["① 設定", "② 実験計画・Y入力", "③ 解析",
                    "④ 次の実験へ" if screening else "④ 最適条件・工程窓", "⑤ 確認実験"])
    with tabs[0]:
        tab_settings(ss)
    with tabs[1]:
        tab_design(ss)
    with tabs[2]:
        tab_analysis(ss)
    with tabs[3]:
        tab_optimum(ss)
    with tabs[4]:
        tab_confirm(ss)
    sidebar_save(ss)


if __name__ == "__main__":
    main()
