"""Excelのプロジェクトファイル：新規作成 → 計画 → Y入力 → 解析 の一連と、v1.1 からの取り込み。"""
import os
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from doe_rsm import analyze
from doe_rsm.__main__ import main
from doe_rsm.excel_io import import_excel_v1, read_project, save_project, write_design, write_results, write_template

FIXTURES = Path(__file__).parent / "fixtures"


def _fill_project(path, design_label="面心CCD", goal="最大化", lsl=48.5, n_center=6, factors=None):
    wb = load_workbook(path)
    ws = wb["Project"]
    vals = {"テーマ": "乾燥条件", "目的変数": "剥離強度", "単位": "N/25mm", "最適化の方向": goal,
            "下限規格 LSL": lsl, "計画の種類": design_label, "中心点の数": n_center, "乱数シード": 42}
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value in vals:
            ws.cell(r, 2, vals[ws.cell(r, 1).value])
    wf = wb["Factors"]
    factors = factors or [("炉温", "℃", 55, 75, None, 80, "○"), ("乾燥時間", "min", 20, 40, 10, None, None),
                          ("循環風速", "m/s", 1, 3, 0.5, 4, None)]
    for i, row in enumerate(factors):
        for j, v in enumerate(row, start=2):
            wf.cell(4 + i, j, v)
    wb.save(path)


def _enter_y(path, func):
    wb = load_workbook(path)
    ws = wb["Design"]
    header = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    cx = header.index("X1(コード)") + 1
    k = sum(1 for h in header if str(h).endswith("(コード)"))
    cy = next(i for i, h in enumerate(header, start=1) if str(h).startswith("Y:"))
    for r in range(4, ws.max_row + 1):
        x = np.array([ws.cell(r, cx + j).value for j in range(k)], float)
        ws.cell(r, cy, func(x))
    wb.save(path)


def test_full_rsm_workflow_reproduces_excel_v1_1(tmp_path):
    """v1.1 の3因子データを新規プロジェクトとして入力し直すと、Excelと同じ推奨条件になる。"""
    import json
    data = json.loads((FIXTURES / "excel_v1_1_k3.json").read_text(encoding="utf-8"))
    y_by_x = {tuple(r["x"]): [] for r in data["runs"]}
    for r in data["runs"]:
        y_by_x[tuple(r["x"])].append(r["y"])
    path = tmp_path / "p.xlsx"
    write_template(str(path), 3)
    _fill_project(path)
    project = read_project(str(path))
    assert project.factors[0].hard_to_change and project.factors[1].safety_low == 10
    d = write_design(str(path), project)
    assert d.n_runs == 20 and sorted(d.run_order) == list(range(1, 21))
    _enter_y(path, lambda x: y_by_x[tuple(x)].pop())
    project = read_project(str(path))
    assert project.n_missing == 0
    res = analyze(project)
    write_results(str(path), project, res)
    exp = data["expected"]
    assert res.fit.sse == pytest.approx(exp["stats"]["sse"], rel=1e-9)
    np.testing.assert_allclose(res.optimization.robust.x, exp["optimization"]["robust"]["x"])
    wb = load_workbook(path)
    for name in ("Project", "Factors", "Design", "Model", "ANOVA", "Gates", "Canonical", "Optimization",
                 "OperatingWindow", "Residuals", "History"):
        assert name in wb.sheetnames
    assert wb.sheetnames.index("Design") == 2 and wb.sheetnames[-1] == "History"
    actions = [wb["History"].cell(r, 2).value for r in range(2, wb["History"].max_row + 1)]
    assert actions[0] == "作成" and "計画" in actions and actions[-1] == "解析"


def test_design_is_reproducible_and_protected(tmp_path):
    path = tmp_path / "p.xlsx"
    write_template(str(path), 3)
    _fill_project(path)
    order1 = write_design(str(path), read_project(str(path))).run_order.copy()
    _enter_y(path, lambda x: 50.0)
    with pytest.raises(ValueError, match="入力済み"):
        write_design(str(path), read_project(str(path)))
    order2 = write_design(str(path), read_project(str(path)), force=True).run_order
    np.testing.assert_array_equal(order1, order2)   # 同じシードなら同じ順序


def test_model_reduction_from_model_sheet(tmp_path):
    path = tmp_path / "p.xlsx"
    write_template(str(path), 3)
    _fill_project(path)
    write_design(str(path), read_project(str(path)))
    rng = np.random.default_rng(5)
    _enter_y(path, lambda x: 50 + 2 * x[0] + x[1] - 2 * x[0] ** 2 - x[1] ** 2 + rng.normal(0, 0.3))
    project = read_project(str(path))
    write_results(str(path), project, analyze(project))
    wb = load_workbook(path)
    ws = wb["Model"]
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 1).value == "X2·X3":
            ws.cell(r, 2, 0)
    wb.save(path)
    project = read_project(str(path))
    assert project.used_terms[6] == 0
    res = analyze(project)
    assert res.fit.coef[6] == 0 and res.fit.n_params == 9


def test_screening_workflow_writes_effects_and_path(tmp_path):
    path = tmp_path / "s.xlsx"
    write_template(str(path), 5)
    factors = [(f"因子{i + 1}", "", 0, 10, None, None, None) for i in range(5)]
    _fill_project(path, design_label="Plackett-Burman", n_center=3, factors=factors)
    d = write_design(str(path), read_project(str(path)))
    assert d.n_runs == 8 + 3
    rng = np.random.default_rng(2)
    _enter_y(path, lambda x: 40 + 3 * x[0] - 2 * x[2] + rng.normal(0, 0.3))
    project = read_project(str(path))
    res = analyze(project)
    assert res.active == ["因子1", "因子3"]
    write_results(str(path), project, res)
    wb = load_workbook(path)
    assert "Effects" in wb.sheetnames and "SteepestPath" in wb.sheetnames and "Model" not in wb.sheetnames


@pytest.mark.parametrize("k", [2, 3, 4])
def test_import_excel_v1_1(tmp_path, k):
    # 元のブックはリポジトリに含めていない。DOE_RSM_V1_DIR に置き場所を指定したときだけ実行する
    src = os.environ.get("DOE_RSM_V1_DIR")
    hits = sorted(Path(src).glob(f"*RSM_{k}*v1.1.xlsx")) if src else []
    if not hits:
        pytest.skip("RSMツール v1.1 のブックがない環境")
    project = import_excel_v1(str(hits[0]))
    dest = tmp_path / "imp.xlsx"
    save_project(str(dest), project)
    again = read_project(str(dest))
    np.testing.assert_array_equal(again.design.run_order, project.design.run_order)
    np.testing.assert_allclose(again.y, project.y)
    assert main(["analyze", str(dest)]) == 0


def test_cli_new_design_and_errors(tmp_path, capsys):
    path = tmp_path / "c.xlsx"
    assert main(["new", str(path), "--factors", "2"]) == 0
    assert main(["design", str(path)]) == 1          # LSL・因子の範囲が未入力
    assert "エラー" in capsys.readouterr().err
    _fill_project(path, factors=[("炉温", "℃", 55, 75, None, None, None), ("時間", "min", 20, 40, None, None, None)])
    assert main(["design", str(path)]) == 0
    assert main(["analyze", str(path)]) == 1         # Y未入力
    assert "未入力" in capsys.readouterr().err
