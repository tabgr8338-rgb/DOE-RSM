"""RSMツール（Excel v1.1）から、入力値と計算結果をテスト用のJSONに書き出す。

使い方:
    python tools/extract_excel_fixture.py RSM_3因子_v1.1.xlsx tests/fixtures/excel_v1_1_k3.json

Excelで一度保存されたファイル（数式の計算結果がキャッシュされているもの）を対象にする。
"""
import json
import sys

import openpyxl


def find_row(ws, text, col=1, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip() == text:
            return r
    raise KeyError(f"{ws.title}: '{text}' が見つからない")


def num(v):
    return None if v in (None, "") else float(v)


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    s1 = wb["01_因子設定"]
    factors = []
    r = 5
    while isinstance(s1.cell(r, 1).value, (int, float)):
        factors.append({
            "name": s1.cell(r, 2).value,
            "unit": s1.cell(r, 4).value,
            "low": num(s1.cell(r, 5).value),
            "high": num(s1.cell(r, 6).value),
            "safety_low": num(s1.cell(r, 9).value),
            "safety_high": num(s1.cell(r, 10).value),
            "hard_to_change": s1.cell(r, 11).value == "○",
        })
        r += 1
    k = len(factors)

    def setting(label):
        return s1.cell(find_row(s1, label), 2).value

    response = {
        "name": setting("名称"),
        "unit": setting("単位"),
        "goal": {"最大化": "maximize", "最小化": "minimize", "目標値": "target"}[setting("最適化の方向")],
        "lsl": num(setting("下限規格 LSL")),
        "usl": num(setting("上限規格 USL")),
        "target": num(setting("目標値 T")),
        "confidence": num(setting("予測区間の信頼水準")),
    }
    design_type = setting("計画の種類")
    n_center = int(setting("中心点の数"))

    # 02: 標準順の計画（コード値）と、03から引いたY
    s2 = wb["02_実験計画"]
    hdr = {s2.cell(4, c).value: c for c in range(1, s2.max_column + 1) if s2.cell(4, c).value}
    y_col = next(c for h, c in hdr.items() if str(h).startswith("Y"))
    runs = []
    for r in range(5, s2.max_row + 1):
        if s2.cell(r, 1).value is None:
            break
        if s2.cell(r, 3).value != 1:
            continue
        runs.append({
            "std_order": int(s2.cell(r, 1).value),
            "kind": s2.cell(r, 2).value,
            "x": [float(s2.cell(r, 4 + i).value) for i in range(k)],
            "run_order": int(s2.cell(r, hdr["実験順序"]).value),
            "y": float(s2.cell(r, y_col).value),
        })

    # 04: 係数・ANOVA・統計量
    s4 = wb["04_回帰ANOVA"]
    terms = []
    r = 6
    while s4.cell(r, 1).value and not str(s4.cell(r, 1).value).startswith("有意"):
        terms.append({
            "term": s4.cell(r, 1).value,
            "used": int(s4.cell(r, 2).value),
            "coef": num(s4.cell(r, 3).value),
            "se": num(s4.cell(r, 4).value),
            "t": num(s4.cell(r, 5).value),
            "p": num(s4.cell(r, 6).value),
        })
        r += 1
    stats = {}
    for label, key in [("SSE", "sse"), ("SST", "sst"), ("SSR", "ssr"), ("MSE", "mse"),
                       ("s（残差標準偏差）", "s"), ("R²", "r2"), ("調整済みR²", "r2_adj"),
                       ("PRESS", "press"), ("予測R²", "r2_pred"),
                       ("t値（工程窓・最適化用）", "t_window"), ("t値（両側・確認実験用）", "t_confirm")]:
        stats[key] = num(s4.cell(find_row(s4, label, col=10), 11).value)
    anova = {}
    for label, key in [("回帰", "regression"), ("残差", "residual"), ("　適合度の欠如", "lack_of_fit"),
                       ("　純誤差", "pure_error")]:
        rr = next(i for i in range(20, 45) if s4.cell(i, 1).value == label)
        anova[key] = {"df": num(s4.cell(rr, 2).value), "ss": num(s4.cell(rr, 3).value),
                      "ms": num(s4.cell(rr, 4).value), "f": num(s4.cell(rr, 5).value),
                      "p": num(s4.cell(rr, 6).value)}

    # 05: スチューデント化残差（標準順）
    s5 = wb["05_残差診断"]
    studentized = {}
    for r in range(5, s5.max_row + 1):
        if s5.cell(r, 3).value == 1:
            studentized[int(s5.cell(r, 1).value)] = num(s5.cell(r, 9).value)

    # 06: 正準解析
    s6 = wb["06_正準解析"]
    stationary = [num(s6.cell(13 + i, 2).value) for i in range(k)]
    eig = [num(s6.cell(23 + i, 2).value) for i in range(k)]
    canonical = {
        "stationary_point": stationary,
        "y_at_stationary": num(s6.cell(18, 2).value),
        "distance": num(s6.cell(19, 2).value),
        "in_region": s6.cell(20, 2).value == "領域内",
        "eigenvalues": eig,
        "kind": s6.cell(28, 2).value,
    }

    # 07: 最適化
    s7 = wb["07_最適化"]
    pred_row = find_row(s7, "予測Y")

    def opt(col):
        return {
            "x": [num(s7.cell(6 + i, col).value) for i in range(k)],
            "pred": num(s7.cell(pred_row, col).value),
            "lower": num(s7.cell(pred_row + 1, col).value),
            "upper": num(s7.cell(pred_row + 2, col).value),
            "margin": num(s7.cell(pred_row + 3, col).value),
        }
    top_row = find_row(s7, "順位") + 1
    top5 = []
    for i in range(5):
        r = top_row + i
        top5.append({
            "x_real": [num(s7.cell(r, 3 + j).value) for j in range(k)],
            "pred": num(s7.cell(r, 3 + k).value),
            "margin": num(s7.cell(r, 6 + k).value),
        })
    optimization = {"robust": opt(2), "point": opt(3), "robust_top5": top5}

    # 08: 1因子窓と組合せチェック
    s8 = wb["08_工程窓"]
    r1 = next(r for r in range(1, s8.max_row + 1)
              if str(s8.cell(r, 1).value or "").startswith("③")) + 2
    one_factor = []
    for i in range(k):
        lo, hi = s8.cell(r1 + i, 3).value, s8.cell(r1 + i, 4).value
        one_factor.append({"low": lo if isinstance(lo, str) else num(lo),
                           "high": hi if isinstance(hi, str) else num(hi)})
    combo = {
        "n_points": num(s8.cell(find_row(s8, "評価点数"), 3).value),
        "n_outside": num(s8.cell(find_row(s8, "安全限界・領域外の点"), 3).value),
        "n_fail": num(s8.cell(find_row(s8, "規格外のおそれがある点"), 3).value),
        "min_margin": num(s8.cell(find_row(s8, "最小の余裕"), 3).value),
        "worst_real": [num(s8.cell(find_row(s8, "最悪の組合せ（実値）"), 3 + j).value) for j in range(k)],
    }

    # 09: 確認実験の1行目（ロバスト最適点での両側予測区間）
    s9 = wb["09_確認実験"]
    confirm = {"pred": num(s9.cell(5, k + 3).value), "pi_low": num(s9.cell(5, k + 4).value),
               "pi_high": num(s9.cell(5, k + 5).value)}

    return {
        "source": path.split("/")[-1],
        "factors": factors,
        "response": response,
        "design_type": design_type,
        "n_center": n_center,
        "runs": runs,
        "expected": {
            "terms": terms, "stats": stats, "anova": anova,
            "studentized_residuals": studentized, "canonical": canonical,
            "optimization": optimization,
            "one_factor_window": one_factor, "combination_check": combo,
            "confirmation_first_row": confirm,
        },
    }


if __name__ == "__main__":
    data = extract(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"wrote {sys.argv[2]}: k={len(data['factors'])}, n={len(data['runs'])}")
